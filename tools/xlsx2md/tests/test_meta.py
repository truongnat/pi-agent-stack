import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color

from xlsx2md.cli import main
from xlsx2md.colors import apply_tint, mark_font, palette_from_theme, resolve_color
from xlsx2md.meta import build_meta
from xlsx2md.tables import process_workbook

_THEME = """<?xml version="1.0"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <a:themeElements>
    <a:clrScheme name="Office">
      <a:dk1><a:srgbClr val="000000"/></a:dk1>
      <a:lt1><a:srgbClr val="FFFFFF"/></a:lt1>
      <a:dk2><a:srgbClr val="1F497D"/></a:dk2>
      <a:lt2><a:srgbClr val="EEECE1"/></a:lt2>
      <a:accent1><a:srgbClr val="4F81BD"/></a:accent1>
    </a:clrScheme>
  </a:themeElements>
</a:theme>
"""


def test_excel_theme_index_is_not_xml_order():
    palette = palette_from_theme(_THEME)
    assert palette[0] == "FFFFFF"
    assert palette[1] == "000000"
    assert palette[4] == "4F81BD"


def test_default_black_font_is_not_a_mark_color():
    palette = palette_from_theme(_THEME)

    class Theme:
        type = "theme"
        theme = 1
        tint = 0

    assert mark_font(resolve_color(Theme(), palette)) is None


def test_indexed_red_and_tint():
    palette = palette_from_theme(_THEME)

    class Indexed:
        type = "indexed"
        indexed = 2
        tint = 0

    assert resolve_color(Indexed(), palette) == "FF0000"
    assert apply_tint("0000FF", 0) == "0000FF"
    light = apply_tint("0000FF", 0.5)
    assert light != "0000FF"
    assert int(light[4:], 16) == 255


def _items_book(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Items"
    fill = PatternFill("solid", fgColor="CCFFCC")
    ws["A1"] = "Items"
    ws["C1"] = "System"
    ws["C1"].fill = fill
    ws["D1"] = "Demo"
    ws["A2"] = "No."
    ws["A2"].fill = fill
    ws["C2"] = "Name"
    ws["C2"].fill = fill
    ws["C2"].comment = Comment("======\nID#x\nAuthor (2026-04-08 07:38:01)\nheader hint", "a")
    ws["E2"] = "Digits"
    ws["E2"].fill = fill
    ws.merge_cells("E2:G2")
    ws["E3"] = "int"
    ws["E3"].fill = fill
    ws["G3"] = "frac"
    ws["G3"].fill = fill
    ws["A4"] = 1
    ws["C4"] = "Clear"
    ws["E4"] = 10
    ws["G4"] = "-"
    ws["C5"] = "Old"
    ws["C5"].font = Font(color="FF0000", strike=True)
    wb.save(path)


def test_meta_maps_fill_comment_strike_and_merge(tmp_path: Path):
    from openpyxl import load_workbook

    path = tmp_path / "sample.xlsx"
    _items_book(path)
    wb = load_workbook(path, data_only=True)
    model = process_workbook(wb)
    document = build_meta(wb, model, source_name=path.name)
    wb.close()

    marks = {(item["sheet"], item["addr"]): item for item in document["marks"]}
    name = marks[("Items", "C2")]
    assert name["fill"] == "CCFFCC"
    assert name["comment"] == "header hint"
    assert name["loc"]["role"] == "header"
    assert name["loc"]["header"] == "Name"
    assert "header hint" not in name["text"]

    old = marks[("Items", "C5")]
    assert old["text"] == "Old"
    assert old["font"] == "FF0000"
    assert old["strike"] is True
    assert old["loc"] == {
        "section": "Items",
        "role": "cell",
        "header": "Name",
        "row": 2,
        "text": "~~Old~~",
    }

    merge = next(item for item in document["merges"] if item["sheet"] == "Items" and item["range"] == "E2:G2")
    assert merge["origin"] == "E2"
    assert merge["text"] == "Digits"
    assert merge["rows"] == 1
    assert merge["cols"] == 3
    assert merge["fill"] == "CCFFCC"
    assert "CCFFCC" in document["colors"]["fill"]
    assert "FF0000" in document["colors"]["font"]


def test_theme_fill_is_exposed_and_default_font_is_not(tmp_path: Path):
    from openpyxl import load_workbook

    src = tmp_path / "theme.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet"
    ws["A1"] = "plain"
    ws["A1"].font = Font(color=Color(theme=1, tint=0.0))
    ws["A1"].fill = PatternFill(fill_type="solid", fgColor=Color(theme=4, tint=0.0))
    ws["B1"] = "note"
    ws["B1"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    ws["C1"] = "red"
    ws["C1"].font = Font(color=Color(indexed=2))
    ws["D1"].fill = PatternFill(fill_type="solid", fgColor="CCFFCC")
    ws.merge_cells("E1:G1")
    ws["E1"].fill = PatternFill(fill_type="solid", fgColor="00B0F0")
    wb.save(src)

    loaded = load_workbook(src, data_only=True)
    document = build_meta(loaded, process_workbook(loaded), source_name=src.name)
    loaded.close()
    marks = {(item["sheet"], item["addr"]): item for item in document["marks"]}
    assert marks[("Sheet", "A1")]["fill"] == "4F81BD"
    assert "font" not in marks[("Sheet", "A1")]
    assert ("Sheet", "B1") not in marks
    assert marks[("Sheet", "C1")]["font"] == "FF0000"
    assert ("Sheet", "D1") not in marks
    assert ("Sheet", "E1") not in marks
    shaded = next(item for item in document["merges"] if item["range"] == "E1:G1")
    assert shaded["fill"] == "00B0F0"
    assert "text" not in shaded
    assert "00B0F0" in document["colors"]["fill"]


def _book(path: Path, *, old: bool) -> None:
    wb = Workbook()
    data = wb.active
    assert data is not None
    data.title = "Data"
    data["A1"] = "Name"
    data["A1"].fill = PatternFill("solid", fgColor="CCFFCC")
    if old:
        data.merge_cells("A1:B1")
    data["A2"] = "Old" if old else "New"
    data["A2"].font = Font(color="FF0000", strike=True) if old else Font(color="0000FF")
    data["B2"] = "keep"
    data["B2"].comment = Comment("keep" if old else "drop", "a")
    if not old:
        data["C2"] = "Extra"
        data["C2"].fill = PatternFill("solid", fgColor="FFFF00")
        added = wb.create_sheet("Added")
        added["A1"] = "y"
    else:
        gone = wb.create_sheet("Gone")
        gone["A1"] = "x"
    wb.save(path)


def test_diff_reports_text_color_comment_merge_and_sheets(tmp_path: Path):
    base = tmp_path / "old.xlsx"
    other = tmp_path / "new.xlsx"
    _book(base, old=True)
    _book(other, old=False)
    diff_md = tmp_path / "diff.md"
    diff_json = tmp_path / "diff.json"
    assert main(["diff", str(base), str(other), "-o", str(diff_md), "--meta", str(diff_json), "-q"]) == 0
    text = diff_md.read_text(encoding="utf-8")
    document = json.loads(diff_json.read_text(encoding="utf-8"))
    kinds = {(item["sheet"], item.get("addr") or item.get("range"), item["kind"]) for item in document["changes"]}
    assert ("Gone", None, "sheet_removed") in kinds
    assert ("Gone", "A1", "removed") in kinds
    assert ("Added", None, "sheet_added") in kinds
    assert ("Added", "A1", "added") in kinds
    assert ("Data", "A2", "changed") in kinds
    assert ("Data", "B2", "changed") in kinds
    assert ("Data", "C2", "added") in kinds
    assert ("Data", "A1:B1", "merge_removed") in kinds
    assert "A2 changed" in text
    assert "FF0000" in text and "0000FF" in text
    changed = next(item for item in document["changes"] if item.get("addr") == "A2")
    assert changed["fields"]["text"] == ["Old", "New"]
    assert changed["fields"]["strike"] == [True, False]
    assert changed["loc"]["section"] == "Data"
    assert "No changes." not in text


def test_cli_writes_meta_sidecar(tmp_path: Path):
    src = tmp_path / "tiny.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Main"
    ws["A1"] = "Label"
    ws["A1"].fill = PatternFill("solid", fgColor="CCFFCC")
    ws["B1"] = "Value"
    ws.merge_cells("B1:C1")
    wb.save(src)
    out = tmp_path / "tiny.md"
    meta = tmp_path / "tiny.meta.json"
    assert main([str(src), "-o", str(out), "--meta", str(meta), "-q"]) == 0
    document = json.loads(meta.read_text(encoding="utf-8"))
    assert document["source"] == "tiny.xlsx"
    assert any(item["addr"] == "A1" and item["fill"] == "CCFFCC" for item in document["marks"])
    assert any(item["range"] == "B1:C1" for item in document["merges"])
    assert "Value" in out.read_text(encoding="utf-8")
