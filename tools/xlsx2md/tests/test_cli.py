from pathlib import Path

from xlsx2md.cli import main


def test_cli_writes_output(tmp_path: Path, capsys):
    from openpyxl import Workbook

    src = tmp_path / "tiny.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "マスタ"
    ws["A1"] = "種類"
    ws["B1"] = "データ型"
    ws["A2"] = "TextBox"
    ws["B2"] = "string"
    wb.save(src)

    out = tmp_path / "tiny.md"
    assert main([str(src), "-o", str(out), "-q"]) == 0
    text = out.read_text(encoding="utf-8")
    assert "## マスタ" in text
    assert "TextBox" in text
    assert "string" in text
