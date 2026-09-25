from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from xlsx2md.cells import comment_body, format_value
from xlsx2md.convert import convert_path
from xlsx2md.tables import process_workbook


def _fill(hex_rgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_rgb)


def test_format_value_int_and_date():
    assert format_value(1.0) == "1"
    assert format_value(datetime(2026, 9, 4)) == "2026-09-04"
    assert format_value("a\nb") == "a / b"


def test_comment_body_strips_gsheets_wrapper():
    raw = "======\nID#AAAB\nFujitsuka Tomohiko    (2026-04-08 07:38:01)\n右、左、中央"
    assert comment_body(raw) == "右、左、中央"


def _sample_workbook(path: Path) -> Path:
    wb = Workbook()
    cover = wb.active
    assert cover is not None
    cover.title = "Cover"
    cover["A1"] = "Spec"
    cover["A2"] = "System"
    cover["B2"] = "Demo"
    cover.merge_cells("B2:F2")
    cover["A3"] = "ID"
    cover["B3"] = "X1"

    items = wb.create_sheet("Items")
    items["A1"] = "Items"
    items["C1"] = "System"
    items["C1"].fill = _fill("CCFFCC")
    items["D1"] = "Demo"
    items["A2"] = "No."
    items["A2"].fill = _fill("CCFFCC")
    items["C2"] = "Name"
    items["C2"].fill = _fill("CCFFCC")
    items["E2"] = "Digits"
    items["E2"].fill = _fill("CCFFCC")
    items.merge_cells("E2:G2")
    items["E3"] = "int"
    items["E3"].fill = _fill("CCFFCC")
    items["G3"] = "frac"
    items["G3"].fill = _fill("CCFFCC")
    items["A4"] = 1.0
    items["C4"] = "Clear"
    items["E4"] = 10
    items["G4"] = "-"
    items["C5"] = "Old"
    items["C5"].font = Font(color="FF0000", strike=True)
    items["C2"].comment = Comment("======\nID#x\nAuthor (2026-04-08 07:38:01)\nheader hint", "a")

    state = wb.create_sheet("State")
    state.merge_cells("E1:H1")
    state["A1"] = "No."
    state["A1"].fill = _fill("DDDDDD")
    state["C1"] = "Name"
    state["C1"].fill = _fill("DDDDDD")
    state["E1"] = "Mode"
    state["E1"].fill = _fill("DDDDDD")
    state["A2"] = 1
    state["C2"] = "Clear"
    state["G2"] = "on"

    items = wb.create_sheet("WideBanner")
    items["A1"] = "WideBanner"
    items["A1"].fill = _fill("CCFFCC")
    items["C1"] = "System"
    items["C1"].fill = _fill("CCFFCC")
    items["E1"] = "DemoApp"
    items["G1"] = "Author"
    items["G1"].fill = _fill("CCFFCC")
    items["I1"] = "Team"
    items["A3"] = "No."
    items["A3"].fill = _fill("CCFFCC")
    items["B3"] = "Name"
    items["B3"].fill = _fill("CCFFCC")
    items["C3"] = "Kind"
    items["C3"].fill = _fill("CCFFCC")
    items["A4"] = 1
    items["B4"] = "Alpha"
    items["C4"] = "Box"
    items["A5"] = 2
    items["B5"] = "Beta"
    items["C5"] = "Btn"

    outline = wb.create_sheet("Notes")
    outline["A1"] = "Init"
    outline["A1"].fill = _fill("CCFFCC")
    outline["C3"] = "1-1."
    outline["G3"] = "Boot screen"
    outline["G4"] = "Set defaults"
    outline["A6"] = "Clear"
    outline["A6"].fill = _fill("CCFFCC")
    outline["C8"] = "Ask"
    outline["I8"] = "Then reset"

    grouped = wb.create_sheet("Grouped")
    grouped["A1"] = "No."
    grouped["A1"].fill = _fill("CCFFCC")
    grouped["B1"] = "Name"
    grouped["B1"].fill = _fill("CCFFCC")
    grouped["C1"] = "Kind"
    grouped["C1"].fill = _fill("CCFFCC")
    grouped["D1"] = "Note"
    grouped["D1"].fill = _fill("CCFFCC")
    grouped["A2"] = "Alpha group"
    grouped["A3"] = 1
    grouped["B3"] = "One"
    grouped["C3"] = "Box"
    grouped["D3"] = "-"
    grouped["A4"] = "Beta group"
    grouped["A5"] = 2
    grouped["B5"] = "Two"
    grouped["C5"] = "Btn"
    grouped["D5"] = "-"

    joins = wb.create_sheet("Joins")
    joins["A1"] = "Process"
    joins["A1"].fill = _fill("CCFFCC")
    for offset, (left, right) in enumerate(
        [("1-1.", "Boot"), ("1-2.", "Load"), ("1-3.", "Draw"), ("1-4.", "Bind"), ("1-5.", "Show")]
    ):
        joins.cell(3 + offset, 3, left)
        joins.cell(3 + offset, 7, right)
    joins["A10"] = "No."
    joins["B10"] = "Join"
    joins["C10"] = "On"
    joins["D10"] = "Note"
    joins["A11"] = 1
    joins["B11"] = "INNER"
    joins["C11"] = "t1=t2"
    joins["A12"] = 2
    joins["B12"] = "LEFT"
    joins["C12"] = "t3=t4"
    joins["D12"] = "ok"
    joins["A14"] = "Wrap"
    joins["A14"].fill = _fill("CCFFCC")

    wb.save(path)
    return path


def test_densify_unmerge_and_header_flatten(tmp_path: Path):
    from openpyxl import load_workbook

    xlsx = _sample_workbook(tmp_path / "sample.xlsx")
    wb = load_workbook(xlsx, data_only=True)
    model = process_workbook(wb)
    wb.close()
    by_name = {table.name: table for table in model.tables}

    items = by_name["Items"]
    assert items.headers[0] == "No."
    assert "Digits/int" in items.headers
    assert "Digits/frac" in items.headers
    assert items.rows[0][0] == "1"
    assert any("~~Old~~" in "".join(row) for row in items.rows)
    assert all("header hint" not in "".join(row) for row in items.rows)

    state = by_name["State"]
    assert state.headers.count("Mode") == 1
    assert any("on" in row for row in state.rows)


def test_typical_width_ignores_wide_banner(tmp_path: Path):
    from openpyxl import load_workbook

    xlsx = _sample_workbook(tmp_path / "sample.xlsx")
    wb = load_workbook(xlsx, data_only=True)
    model = process_workbook(wb)
    wb.close()
    table = next(item for item in model.tables if item.name == "WideBanner")
    assert table.headers == ["No.", "Name", "Kind"]
    assert table.rows[0] == ["1", "Alpha", "Box"]
    assert model.meta.get("System") == "Demo"


def test_sparse_sheet_renders_as_outline(tmp_path: Path):
    markdown, model = convert_path(_sample_workbook(tmp_path / "sample.xlsx"))
    notes = next(item for item in model.tables if item.name == "Notes")
    assert notes.sparse
    assert "Boot screen" in markdown
    assert "**Init**" in markdown
    assert "**Clear**" in markdown
    assert "**Clear**\n  Ask" in markdown
    assert "Notes_2" not in markdown
    cover = markdown.split("## Cover", 1)[1].split("##", 1)[0]
    assert "Spec" in cover
    assert "System" not in cover
    assert "| c1 |" not in markdown


def test_banner_rows_split_dense_tables(tmp_path: Path):
    markdown, model = convert_path(_sample_workbook(tmp_path / "sample.xlsx"))
    grouped = next(item for item in model.tables if item.name == "Grouped")
    assert not grouped.sparse
    assert "**Alpha group**" in markdown
    assert "**Beta group**" in markdown
    assert "| Alpha group |" not in markdown


def test_one_cell_data_rows_are_not_banners(tmp_path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    src = tmp_path / "master.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Master"
    fill = PatternFill("solid", fgColor="CCFFCC")
    ws["A1"] = "Kind"
    ws["A1"].fill = fill
    ws["B1"] = "Type"
    ws["B1"].fill = fill
    ws["C1"] = "Align"
    ws["C1"].fill = fill
    ws["A2"] = "TextBox"
    ws["B2"] = "string"
    ws["C2"] = "left"
    ws["A3"] = "Button"
    ws["A4"] = "IconButton"
    ws["A5"] = "Link"
    wb.save(src)

    markdown, model = convert_path(src)
    table = next(item for item in model.tables if item.name == "Master")
    assert not table.sparse
    assert any(row[0] == "Button" for row in table.rows)
    assert "**Button**" not in markdown
    assert "IconButton" in markdown


def test_outline_numbered_block_becomes_table(tmp_path: Path):
    markdown, model = convert_path(_sample_workbook(tmp_path / "sample.xlsx"))
    joins = next(item for item in model.tables if item.name == "Joins")
    assert joins.sparse
    assert "| No. | Join | On | Note |" in markdown
    assert "| 1 | INNER | t1=t2 |" in markdown
    assert "LEFT" in markdown
    assert "**Process**" in markdown
    assert "**Wrap**" in markdown


def test_markitdown_html_path_has_no_nan(tmp_path: Path):
    xlsx = _sample_workbook(tmp_path / "sample.xlsx")
    markdown, model = convert_path(xlsx)
    assert model is not None
    assert "Unnamed:" not in markdown
    assert "| NaN |" not in markdown
    assert "## Items" in markdown
    assert "Digits/int" in markdown
    assert "~~Old~~" in markdown


def test_keep_xlsx_is_pandas_friendly(tmp_path: Path):
    import pandas as pd

    xlsx = _sample_workbook(tmp_path / "sample.xlsx")
    dense = tmp_path / "dense.xlsx"
    convert_path(xlsx, keep_xlsx=dense)
    frame = pd.read_excel(dense, sheet_name="Items")
    assert "Unnamed: 1" not in list(frame.columns)
    assert "Name" in list(frame.columns)
