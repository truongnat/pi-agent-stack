"""Turn a sparse, merged worksheet into dense tables or an outline (geometry only)."""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Iterable

from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .cells import Cell, load_sheet
from .colors import has_fill


@dataclass(slots=True)
class Table:
    name: str
    headers: list[str]
    rows: list[list[str]]
    source_sheet: str
    in_cells: int
    intro: list[str] = field(default_factory=list)
    pairs: list[tuple[str, str]] = field(default_factory=list)
    sparse: bool = False
    row_addrs: list[list[str]] = field(default_factory=list)
    header_addrs: list[list[str]] = field(default_factory=list)
    pair_addrs: list[tuple[str, str]] = field(default_factory=list)
    intro_addrs: list[str] = field(default_factory=list)


@dataclass(slots=True)
class WorkbookModel:
    title: str
    meta: dict[str, str] = field(default_factory=dict)
    tables: list[Table] = field(default_factory=list)
    # sheet, addr, meta key — cells lifted into the workbook meta table
    meta_refs: list[tuple[str, str, str]] = field(default_factory=list)


def _addr(cell: Cell) -> str:
    return f"{get_column_letter(cell.col)}{cell.row}"


def _rows_of(cells: dict[tuple[int, int], Cell]) -> list[int]:
    return sorted({row for row, _ in cells})


def _cols_of(cells: dict[tuple[int, int], Cell], rows: Iterable[int]) -> list[int]:
    wanted = set(rows)
    return sorted({col for row, col in cells if row in wanted})


def _row_cells(cells: dict[tuple[int, int], Cell], row: int) -> list[tuple[int, Cell]]:
    return sorted(
        ((col, cells[(row, col)]) for r, col in cells if r == row),
        key=lambda item: item[0],
    )


def _width(cells: dict[tuple[int, int], Cell], row: int) -> int:
    return len(_row_cells(cells, row))


def cell_text(cell: Cell, *, markup: bool) -> str:
    text = cell.text
    if not markup:
        return text
    if cell.strike and text:
        text = f"~~{text}~~"
    if cell.comment:
        text = f"{text} ({cell.comment})" if text else f"({cell.comment})"
    return text


def _unique(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out: list[str] = []
    for name in headers:
        n = seen.get(name, 0) + 1
        seen[name] = n
        out.append(name if n == 1 else f"{name}_{n}")
    return out


def flatten_headers(
    cells: dict[tuple[int, int], Cell],
    header_rows: list[int],
    cols: list[int],
    merge_origin: dict[tuple[int, int], tuple[int, int]],
) -> tuple[list[str], list[list[str]]]:
    headers: list[str] = []
    addrs: list[list[str]] = []
    for col in cols:
        parts: list[str] = []
        used: list[str] = []
        for row in header_rows:
            origin = merge_origin.get((row, col), (row, col))
            cell = cells.get(origin) or cells.get((row, col))
            if cell and cell.text:
                token = cell.text.replace(" / ", "")
                if token not in parts:
                    parts.append(token)
                    used.append(_addr(cell))
        headers.append("/".join(parts) if parts else f"col{col}")
        addrs.append(used)
    return headers, addrs


def _take(row: list[str], indexes: list[int]) -> list[str]:
    return [row[i] if i < len(row) else "" for i in indexes]


def _drop_empty_columns(
    headers: list[str],
    rows: list[list[str]],
    row_addrs: list[list[str]],
    header_addrs: list[list[str]],
) -> tuple[list[str], list[list[str]], list[list[str]], list[list[str]]]:
    if not headers:
        return headers, rows, row_addrs, header_addrs
    keep = [i for i, _ in enumerate(headers) if any(r[i] for r in rows if i < len(r))]
    if not keep or len(keep) == len(headers):
        return headers, rows, row_addrs, header_addrs
    return (
        [headers[i] for i in keep],
        [_take(row, keep) for row in rows],
        [_take(row, keep) for row in row_addrs],
        [header_addrs[i] if i < len(header_addrs) else [] for i in keep],
    )


def _typical_width(cells: dict[tuple[int, int], Cell], rows: list[int]) -> int:
    widths = [_width(cells, row) for row in rows]
    clustered = [width for width in widths if width >= 2]
    if not clustered:
        return max(widths) if widths else 0
    counts = Counter(clustered)
    return max(counts, key=lambda width: (counts[width], -width))


def _fill_ratio(cells: dict[tuple[int, int], Cell], row: int) -> float:
    items = _row_cells(cells, row)
    if not items:
        return 0.0
    return sum(1 for _, cell in items if has_fill(cell.fill_rgb)) / len(items)


def _label_value_pairs(
    items: list[tuple[int, Cell]], *, markup: bool
) -> tuple[list[tuple[str, str]], list[str], list[tuple[str, str]], list[str]]:
    """Filled cell + following unfilled cell → label/value. Leftovers become intro lines."""
    pairs: list[tuple[str, str]] = []
    leftover: list[str] = []
    pair_addrs: list[tuple[str, str]] = []
    leftover_addrs: list[str] = []
    index = 0
    while index < len(items):
        _, left = items[index]
        nxt = items[index + 1][1] if index + 1 < len(items) else None
        if nxt is not None and has_fill(left.fill_rgb) and not has_fill(nxt.fill_rgb) and nxt.text:
            pairs.append((left.text, cell_text(nxt, markup=markup)))
            pair_addrs.append((_addr(left), _addr(nxt)))
            index += 2
            continue
        leftover.append(cell_text(left, markup=markup))
        leftover_addrs.append(_addr(left))
        index += 1
    return pairs, leftover, pair_addrs, leftover_addrs


def _header_band(cells: dict[tuple[int, int], Cell], rows: list[int], typical: int) -> list[int] | None:
    if not rows or typical < 3:
        return None
    low = max(3, int(typical * 0.7))
    high = typical + 2
    wide = [row for row in rows if low <= _width(cells, row) <= high]
    if not wide:
        return None
    filled = [row for row in wide if _fill_ratio(cells, row) >= 0.5]
    candidates = filled or wide

    def follow_score(start: int) -> int:
        need = max(2, int(_width(cells, start) * 0.5))
        n = 0
        for row in rows:
            if row <= start:
                continue
            if n == 0 and _fill_ratio(cells, row) >= 0.5:
                continue
            if low <= _width(cells, row) <= high or _width(cells, row) >= need:
                n += 1
            elif n:
                break
        return n

    start = max(candidates, key=lambda row: (follow_score(row), row))
    if follow_score(start) < 1 and _fill_ratio(cells, start) < 0.5:
        return None
    band = [start]
    idx = rows.index(start)
    while idx + 1 < len(rows) and len(band) < 3:
        nxt = rows[idx + 1]
        nxt_cells = _row_cells(cells, nxt)
        if not nxt_cells:
            break
        if _fill_ratio(cells, nxt) >= 0.5 and 2 <= len(nxt_cells) <= high:
            band.append(nxt)
            idx += 1
            continue
        break
    return band


def _matrix(
    cells: dict[tuple[int, int], Cell],
    rows: list[int],
    cols: list[int],
    *,
    markup: bool,
) -> tuple[list[list[str]], list[list[str]]]:
    out: list[list[str]] = []
    addrs: list[list[str]] = []
    for row in rows:
        line: list[str] = []
        line_addrs: list[str] = []
        for col in cols:
            cell = cells.get((row, col))
            line.append(cell_text(cell, markup=markup) if cell else "")
            line_addrs.append(_addr(cell) if cell else "")
        if any(line):
            out.append(line)
            addrs.append(line_addrs)
    return out, addrs


def _filled(row: list[str]) -> list[tuple[int, str]]:
    return [(index, cell) for index, cell in enumerate(row) if cell]


def _is_numeric(text: str) -> bool:
    return text.replace(".", "", 1).isdigit()


def _drop_trailing_stubs(
    rows: list[list[str]], addrs: list[list[str]]
) -> tuple[list[list[str]], list[list[str]]]:
    def is_stub(row: list[str]) -> bool:
        filled = _filled(row)
        return len(filled) == 1 and filled[0][0] == 0 and _is_numeric(filled[0][1])

    end = len(rows)
    while end > 0 and is_stub(rows[end - 1]):
        end -= 1
    if end >= 2 and end < len(rows):
        return rows[:end], addrs[:end]
    return rows, addrs


def _echoes_meta(row: list[str], meta: dict[str, str]) -> bool:
    filled = [cell for _, cell in _filled(row)]
    if len(filled) == 2 and meta.get(filled[0]) == filled[1]:
        return True
    if len(filled) == 1:
        text = filled[0]
        for key, value in meta.items():
            if text in {f"{key} {value}", f"{key}: {value}", f"{key}：{value}"}:
                return True
    return False


def _is_sparse(headers: list[str], rows: list[list[str]]) -> bool:
    if not rows or len(headers) < 4:
        return False
    filled = sum(1 for row in rows for cell in row if cell)
    occupancy = filled / (len(rows) * len(headers))
    widths = sorted(sum(1 for cell in row if cell) for row in rows)
    median = widths[len(widths) // 2]
    return occupancy < 0.28 or (median <= 3 and len(headers) >= 4)


def _keep_intro(texts: list[str], addrs: list[str], title: str) -> tuple[list[str], list[str]]:
    kept_text: list[str] = []
    kept_addr: list[str] = []
    for text, addr in zip(texts, addrs):
        if text and text != title:
            kept_text.append(text)
            kept_addr.append(addr)
    return kept_text, kept_addr


def process_sheet(ws: Worksheet, *, markup: bool = True) -> list[Table]:
    cells, merge_origin = load_sheet(ws)
    in_cells = len(cells)
    if not cells:
        return []

    rows = _rows_of(cells)
    typical = _typical_width(cells, rows)
    band = _header_band(cells, rows, typical)
    pairs: list[tuple[str, str]] = []
    pair_addrs: list[tuple[str, str]] = []
    intro: list[str] = []
    intro_addrs: list[str] = []
    if band:
        chrome_rows = [row for row in rows if row < band[0]]
        body_rows = [row for row in rows if row > band[-1]]
    else:
        body_rows = list(rows)
        chrome_rows = []
        while body_rows:
            items = _row_cells(cells, body_rows[0])
            row_pairs, leftover, row_pair_addrs, leftover_addrs = _label_value_pairs(items, markup=markup)
            if row_pairs:
                body_rows.pop(0)
                pairs.extend(row_pairs)
                pair_addrs.extend(row_pair_addrs)
                extra, extra_addrs = _keep_intro(leftover, leftover_addrs, ws.title)
                intro.extend(extra)
                intro_addrs.extend(extra_addrs)
                continue
            break

    for row in chrome_rows:
        items = _row_cells(cells, row)
        row_pairs, leftover, row_pair_addrs, leftover_addrs = _label_value_pairs(items, markup=markup)
        pairs.extend(row_pairs)
        pair_addrs.extend(row_pair_addrs)
        extra, extra_addrs = _keep_intro(leftover, leftover_addrs, ws.title)
        intro.extend(extra)
        intro_addrs.extend(extra_addrs)

    cols = _cols_of(cells, (band or []) + body_rows)
    if not body_rows or not cols:
        if intro or pairs:
            return [
                Table(
                    name=ws.title,
                    headers=[],
                    rows=[],
                    source_sheet=ws.title,
                    in_cells=in_cells,
                    intro=intro,
                    pairs=pairs,
                    pair_addrs=pair_addrs,
                    intro_addrs=intro_addrs,
                )
            ]
        return []

    if band:
        headers, header_addrs = flatten_headers(cells, band, cols, merge_origin)
        matrix, row_addrs = _matrix(cells, body_rows, cols, markup=markup)
        headers, matrix, row_addrs, header_addrs = _drop_empty_columns(headers, matrix, row_addrs, header_addrs)
        headers = _unique(headers)
    else:
        matrix, row_addrs = _matrix(cells, body_rows, cols, markup=markup)
        headers = [f"c{i}" for i in range(1, (len(matrix[0]) if matrix else len(cols)) + 1)]
        header_addrs = [[] for _ in headers]
        headers, matrix, row_addrs, header_addrs = _drop_empty_columns(headers, matrix, row_addrs, header_addrs)

    matrix, row_addrs = _drop_trailing_stubs(matrix, row_addrs)
    if not matrix and not intro and not pairs:
        return []
    return [
        Table(
            name=ws.title,
            headers=headers,
            rows=matrix,
            source_sheet=ws.title,
            in_cells=in_cells,
            intro=intro,
            pairs=pairs,
            sparse=band is None or _is_sparse(headers, matrix),
            row_addrs=row_addrs,
            header_addrs=header_addrs,
            pair_addrs=pair_addrs,
            intro_addrs=intro_addrs,
        )
    ]


def process_workbook(wb: Workbook, *, markup: bool = True, sheets: list[str] | None = None) -> WorkbookModel:
    wanted = set(sheets) if sheets else None
    model = WorkbookModel(title=wb.properties.title or "")
    for ws in wb.worksheets:
        if wanted is not None and ws.title not in wanted:
            continue
        model.tables.extend(process_sheet(ws, markup=markup))

    counts: Counter[str] = Counter()
    values: dict[str, str] = {}
    for table in model.tables:
        for key, value in table.pairs:
            counts[key] += 1
            values.setdefault(key, value)
    model.meta = {key: values[key] for key, n in counts.items() if n >= 2}
    if model.meta:
        kept: list[Table] = []
        refs: list[tuple[str, str, str]] = []
        for table in model.tables:
            kept_pairs: list[tuple[str, str]] = []
            kept_pair_addrs: list[tuple[str, str]] = []
            for (key, value), addrs in zip(table.pairs, table.pair_addrs):
                if key in model.meta:
                    refs.append((table.source_sheet, addrs[0], key))
                    refs.append((table.source_sheet, addrs[1], key))
                    continue
                kept_pairs.append((key, value))
                kept_pair_addrs.append(addrs)
            table.pairs = kept_pairs
            table.pair_addrs = kept_pair_addrs
            kept_rows: list[list[str]] = []
            kept_addrs: list[list[str]] = []
            for row, addrs in zip(table.rows, table.row_addrs):
                if _echoes_meta(row, model.meta):
                    filled = _filled(row)
                    if len(filled) == 2 and filled[0][1] in model.meta:
                        key = filled[0][1]
                        label_at = filled[0][0]
                        value_at = filled[1][0]
                        if label_at < len(addrs) and addrs[label_at]:
                            refs.append((table.source_sheet, addrs[label_at], key))
                        if value_at < len(addrs) and addrs[value_at]:
                            refs.append((table.source_sheet, addrs[value_at], key))
                    continue
                kept_rows.append(row)
                kept_addrs.append(addrs)
            table.rows = kept_rows
            table.row_addrs = kept_addrs
            if table.rows or table.intro or table.pairs:
                kept.append(table)
        model.tables = kept
        model.meta_refs = refs

    if not model.title and model.tables:
        model.title = model.tables[0].source_sheet
    return model


def write_tables_xlsx(model: WorkbookModel, path: str) -> None:
    wb = Workbook()
    default = wb.active
    assert default is not None
    first = True
    if model.meta:
        ws = default
        ws.title = "meta"
        ws["A1"] = "key"
        ws["B1"] = "value"
        for index, (key, value) in enumerate(model.meta.items(), start=2):
            ws.cell(index, 1, key)
            ws.cell(index, 2, value)
        first = False
    for table in model.tables:
        if first:
            ws = default
            first = False
        else:
            ws = wb.create_sheet()
        name = re.sub(r"[\\/*?:\[\]]", "-", table.name)[:31] or "Sheet"
        ws.title = name
        for col, header in enumerate(table.headers, start=1):
            ws.cell(1, col, header)
        for r_index, row in enumerate(table.rows, start=2):
            for c_index, value in enumerate(row, start=1):
                ws.cell(r_index, c_index, value or None)
    if first:
        default["A1"] = "(empty)"
    wb.save(path)
