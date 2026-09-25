"""Compare two workbooks by cell text and by the facts Markdown drops."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple

from .colors import theme_palette
from .facts import Fact, Merge, collect_sheet
from .meta import index_locs
from .tables import WorkbookModel, process_workbook

_FIELDS = ("text", "fill", "font", "strike", "bold", "comment")


def _sig(fact: Fact) -> tuple:
    return (fact.text, fact.fill, fact.font, fact.strike, fact.bold, fact.comment)


def _addr_key(addr: str) -> tuple[int, int]:
    origin = addr.split(":", 1)[0]
    try:
        row, col = coordinate_to_tuple(origin)
    except (ValueError, TypeError):
        return 0, 0
    return row, col


def _load(
    path: Path, sheets: list[str] | None
) -> tuple[dict[tuple[str, str], Fact], list[Merge], WorkbookModel, list[str]]:
    wb = load_workbook(path, data_only=True)
    try:
        palette = theme_palette(wb)
        wanted = set(sheets) if sheets else None
        facts: dict[tuple[str, str], Fact] = {}
        merges: list[Merge] = []
        order: list[str] = []
        for ws in wb.worksheets:
            if wanted is not None and ws.title not in wanted:
                continue
            order.append(ws.title)
            sheet_facts, sheet_merges = collect_sheet(ws, palette)
            for fact in sheet_facts.values():
                facts[(fact.sheet, fact.addr)] = fact
            merges.extend(sheet_merges)
        model = process_workbook(wb, sheets=sheets)
        model.title = model.title or path.name
    finally:
        wb.close()
    return facts, merges, model, order


def _show(value: object) -> str:
    if isinstance(value, bool):
        return "true" if value else "false"
    if value is None or value == "":
        return "(none)"
    return "`" + str(value).replace("`", "'") + "`"


def _field(fact: Fact, name: str) -> object:
    return getattr(fact, name)


def compare(
    base: dict[tuple[str, str], Fact],
    other: dict[tuple[str, str], Fact],
    base_merges: list[Merge],
    other_merges: list[Merge],
    *,
    base_order: list[str],
    other_order: list[str],
) -> list[dict]:
    changes: list[dict] = []
    base_sheets = set(base_order)
    other_sheets = set(other_order)
    sheet_order = list(base_order)
    sheet_order.extend(name for name in other_order if name not in base_sheets)
    rank = {name: index for index, name in enumerate(sheet_order)}

    def _rank(name: str) -> int:
        return rank.get(name, 10**6)

    for name in sheet_order:
        if name not in other_sheets and name in base_sheets:
            changes.append({"sheet": name, "kind": "sheet_removed"})
        elif name not in base_sheets and name in other_sheets:
            changes.append({"sheet": name, "kind": "sheet_added"})

    keys = set(base) | set(other)
    for sheet, addr in sorted(keys, key=lambda item: (_rank(item[0]), _addr_key(item[1]))):
        left = base.get((sheet, addr))
        right = other.get((sheet, addr))
        if left is None and right is not None:
            item = {"sheet": sheet, "addr": addr, "kind": "added", **_present(right)}
            changes.append(item)
            continue
        if right is None and left is not None:
            changes.append({"sheet": sheet, "addr": addr, "kind": "removed", **_present(left)})
            continue
        if left is None or right is None or _sig(left) == _sig(right):
            continue
        fields = {}
        for name in _FIELDS:
            old = _field(left, name)
            new = _field(right, name)
            if old != new:
                fields[name] = [old, new]
        if fields:
            changes.append({"sheet": sheet, "addr": addr, "kind": "changed", "fields": fields})

    base_ranges = {(item.sheet, item.range): item for item in base_merges}
    other_ranges = {(item.sheet, item.range): item for item in other_merges}
    for key in sorted(set(base_ranges) | set(other_ranges), key=lambda item: (_rank(item[0]), _addr_key(item[1]))):
        if key in base_ranges and key in other_ranges:
            continue
        merge = other_ranges.get(key) or base_ranges[key]
        kind = "merge_added" if key in other_ranges else "merge_removed"
        item = {"sheet": merge.sheet, "range": merge.range, "kind": kind, "origin": merge.origin}
        if merge.text:
            item["text"] = merge.text
        changes.append(item)
    return changes


def _present(fact: Fact) -> dict:
    data: dict = {}
    if fact.text:
        data["text"] = fact.text
    if fact.fill:
        data["fill"] = fact.fill
    if fact.font:
        data["font"] = fact.font
    if fact.strike:
        data["strike"] = True
    if fact.bold:
        data["bold"] = True
    if fact.comment:
        data["comment"] = fact.comment
    return data


def attach_locs(changes: list[dict], *, base_model: WorkbookModel, other_model: WorkbookModel) -> None:
    base_locs = index_locs(base_model)
    other_locs = index_locs(other_model)
    for change in changes:
        kind = change["kind"]
        locs = other_locs if kind in {"added", "changed", "merge_added"} else base_locs
        addr = change.get("addr") or change.get("origin")
        if not addr:
            continue
        loc = locs.get((change["sheet"], addr))
        if loc is None and kind == "changed":
            loc = base_locs.get((change["sheet"], addr))
        if loc:
            change["loc"] = loc


def render_diff(base_name: str, other_name: str, changes: list[dict]) -> str:
    lines = [f"# {base_name} → {other_name}", ""]
    if not changes:
        lines.append("No changes.")
        return "\n".join(lines) + "\n"
    grouped: dict[str, list[dict]] = {}
    order: list[str] = []
    for change in changes:
        sheet = change["sheet"]
        if sheet not in grouped:
            grouped[sheet] = []
            order.append(sheet)
        grouped[sheet].append(change)
    for sheet in order:
        lines.append(f"## {sheet}")
        for change in grouped[sheet]:
            lines.append(f"- {_line(change)}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def _line(change: dict) -> str:
    kind = change["kind"]
    if kind == "sheet_added":
        return "sheet added"
    if kind == "sheet_removed":
        return "sheet removed"
    if kind in {"merge_added", "merge_removed"}:
        verb = "merge added" if kind == "merge_added" else "merge removed"
        text = f" ({change['text']})" if change.get("text") else ""
        return f"{change['range']} {verb}{text}"
    if kind == "changed":
        parts = []
        for name, (old, new) in change.get("fields", {}).items():
            parts.append(f"{name}: {_show(old)} → {_show(new)}")
        return f"{change['addr']} changed " + "; ".join(parts)
    bits = [f"{change['addr']} {kind}"]
    for name in ("text", "fill", "font", "comment"):
        if change.get(name):
            bits.append(f"{name}={_show(change[name])}")
    if change.get("strike"):
        bits.append("strike=true")
    if change.get("bold"):
        bits.append("bold=true")
    return " ".join(bits)


def diff_paths(
    base: str | Path,
    other: str | Path,
    *,
    sheets: list[str] | None = None,
) -> tuple[str, dict, int]:
    base_path = Path(base)
    other_path = Path(other)
    base_facts, base_merges, base_model, base_order = _load(base_path, sheets)
    other_facts, other_merges, other_model, other_order = _load(other_path, sheets)
    changes = compare(
        base_facts,
        other_facts,
        base_merges,
        other_merges,
        base_order=base_order,
        other_order=other_order,
    )
    attach_locs(changes, base_model=base_model, other_model=other_model)
    markdown = render_diff(base_path.name, other_path.name, changes)
    document = {"base": base_path.name, "other": other_path.name, "changes": changes}
    return markdown, document, len(changes)


def dump_diff(document: dict, path: str | Path) -> None:
    Path(path).write_text(json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
