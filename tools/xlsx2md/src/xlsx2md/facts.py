"""Cell facts Markdown cannot carry: color, strike, bold, comment, merge."""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .cells import _iter_existing_cells, comment_body, format_value
from .colors import mark_fill, mark_font, resolve_color


@dataclass(slots=True)
class Fact:
    sheet: str
    addr: str
    row: int
    col: int
    text: str
    fill: str | None = None
    font: str | None = None
    strike: bool = False
    bold: bool = False
    comment: str | None = None
    merge: str | None = None

    def is_mark(self) -> bool:
        return bool(self.fill or self.font or self.strike or self.bold or self.comment)

    def as_dict(self, *, marks_only: bool = False) -> dict:
        data: dict = {"sheet": self.sheet, "addr": self.addr, "text": self.text}
        if self.fill:
            data["fill"] = self.fill
        if self.font:
            data["font"] = self.font
        if self.strike:
            data["strike"] = True
        if self.bold:
            data["bold"] = True
        if self.comment:
            data["comment"] = self.comment
        if self.merge and (not marks_only or self.is_mark()):
            data["merge"] = self.merge
        return data


@dataclass(slots=True)
class Merge:
    sheet: str
    range: str
    origin: str
    text: str
    rows: int
    cols: int
    fill: str | None = None

    def as_dict(self) -> dict:
        data = {
            "sheet": self.sheet,
            "range": self.range,
            "origin": self.origin,
            "rows": self.rows,
            "cols": self.cols,
        }
        if self.text:
            data["text"] = self.text
        if self.fill:
            data["fill"] = self.fill
        return data


def _fill_of(cell, palette: list[str]) -> str | None:
    fill = cell.fill
    if fill is None or fill.fill_type in (None, "none"):
        return None
    return mark_fill(resolve_color(getattr(fill, "fgColor", None), palette))


def _font_of(cell, palette: list[str]) -> tuple[str | None, bool, bool]:
    font = cell.font
    if font is None:
        return None, False, False
    return mark_font(resolve_color(font.color, palette)), bool(font.strike), bool(font.bold)


def collect_sheet(ws: Worksheet, palette: list[str]) -> tuple[dict[str, Fact], list[Merge]]:
    merge_origin: dict[tuple[int, int], tuple[int, int]] = {}
    merges: list[Merge] = []
    for merged in ws.merged_cells.ranges:
        origin = (merged.min_row, merged.min_col)
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                merge_origin[(row, col)] = origin
        merges.append(
            Merge(
                sheet=ws.title,
                range=merged.coord,
                origin=f"{get_column_letter(merged.min_col)}{merged.min_row}",
                text="",
                rows=merged.max_row - merged.min_row + 1,
                cols=merged.max_col - merged.min_col + 1,
            )
        )

    facts: dict[str, Fact] = {}
    origin_fill: dict[str, str] = {}
    origins = {item.origin for item in merges}
    for xl in _iter_existing_cells(ws):
        origin = merge_origin.get((xl.row, xl.column))
        if origin and origin != (xl.row, xl.column) and xl.value is None and xl.comment is None:
            continue
        text = format_value(xl.value)
        comment = comment_body(xl.comment)
        fill = _fill_of(xl, palette)
        font, strike, bold = _font_of(xl, palette)
        addr = f"{get_column_letter(xl.column)}{xl.row}"
        if fill and addr in origins:
            origin_fill[addr] = fill
        # Empty fill with no other fact is unused-column shading, not a cell to map.
        if not text and not comment and not font and not strike and not bold:
            continue
        facts[addr] = Fact(
            sheet=ws.title,
            addr=addr,
            row=xl.row,
            col=xl.column,
            text=text,
            fill=fill,
            font=font,
            strike=strike,
            bold=bold,
            comment=comment,
        )

    for merge in merges:
        merge.fill = origin_fill.get(merge.origin)
        fact = facts.get(merge.origin)
        if fact is None:
            continue
        merge.text = fact.text
        fact.merge = merge.range
    return facts, merges
