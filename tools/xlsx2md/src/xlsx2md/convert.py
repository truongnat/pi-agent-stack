"""Run markitdown on dense tables (same HTML path as its XLSX converter)."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from markitdown import MarkItDown
from markitdown.converters._html_converter import HtmlConverter
from openpyxl import load_workbook

from .tables import Table, WorkbookModel, process_workbook, write_tables_xlsx

_HTML = HtmlConverter()
_NA_CELL = re.compile(r"(?<=\|)\s*NaN\s*(?=\|)")
_MULTI_BLANK = re.compile(r"\n{3,}")
_STEP = re.compile(r"^\d+([.\-]\d+)*[.)．、]?\s*$")


def _df_to_markdown(headers: list[str], rows: list[list[str]]) -> str:
    frame = pd.DataFrame(rows, columns=headers)
    payload = frame.to_html(index=False, na_rep="", border=0, escape=True)
    markdown = _HTML.convert_string(payload).markdown.strip()
    markdown = _NA_CELL.sub(" ", markdown)
    markdown = markdown.replace("\\|", "｜")
    markdown = markdown.replace("\\_", "_")
    return markdown


def _filled(row: list[str]) -> list[tuple[int, str]]:
    return [(index, cell) for index, cell in enumerate(row) if cell]


def _is_numeric(text: str) -> bool:
    return text.replace(".", "", 1).isdigit()


def _is_bullet(line: str) -> bool:
    return line.startswith(("・", "•", "- ", "* ", "– "))


def _render_intro(lines: list[str]) -> str:
    blocks: list[str] = []
    bullets: list[str] = []

    def flush() -> None:
        if bullets:
            blocks.append("\n".join(bullets))
            bullets.clear()

    for line in lines:
        if _is_bullet(line):
            bullets.append(line)
            continue
        flush()
        if len(line) > 40:
            blocks.append(line)
        else:
            blocks.append(line if line.startswith("**") else f"**{line}**")
    flush()
    return "\n\n".join(blocks)


def _banner_title(row: list[str]) -> str | None:
    filled = _filled(row)
    if len(filled) != 1 or filled[0][0] != 0:
        return None
    text = filled[0][1]
    if len(text) < 2 or _is_numeric(text):
        return None
    return text


def _has_following_data(rows: list[list[str]], start: int) -> bool:
    for row in rows[start + 1 :]:
        if len(_filled(row)) >= 2:
            return True
    return False


def banner_spans(rows: list[list[str]]) -> list[tuple[str | None, int | None, int, int]]:
    """(title, title row index, start, end) slices of body rows. Title rows are outside the slice."""

    def is_banner(index: int) -> bool:
        return _banner_title(rows[index]) is not None and _has_following_data(rows, index)

    if not rows or not any(is_banner(index) for index in range(len(rows))):
        return [(None, None, 0, len(rows))]

    spans: list[tuple[str | None, int | None, int, int]] = []
    title: str | None = None
    title_index: int | None = None
    start = 0
    for index in range(len(rows)):
        if not is_banner(index):
            continue
        if index > start or title is not None:
            spans.append((title, title_index, start, index))
        title = _banner_title(rows[index])
        title_index = index
        start = index + 1
    if start < len(rows) or title is not None:
        spans.append((title, title_index, start, len(rows)))
    return spans


def _split_banner_chunks(rows: list[list[str]]) -> list[tuple[str | None, list[list[str]]]]:
    return [(title, rows[start:end]) for title, _index, start, end in banner_spans(rows)]


def _mini_run_len(rows: list[list[str]], start: int) -> int:
    head = _filled(rows[start])
    if len(head) < 3 or _is_numeric(head[0][1]) or _STEP.match(head[0][1]):
        return 0
    col0 = head[0][0]
    numbered = 0
    n = 1
    widths = [len(head)]
    for row in rows[start + 1 :]:
        filled = _filled(row)
        if len(filled) < 3 or filled[0][0] != col0:
            break
        n += 1
        widths.append(len(filled))
        if _is_numeric(filled[0][1]):
            numbered += 1
    if n < 2:
        return 0
    if numbered >= n - 1:
        return n
    if numbered == 0 and max(widths) - min(widths) <= 1:
        return n
    return 0


def _texts(row: list[str]) -> list[str]:
    return [cell for _, cell in _filled(row)]


def _table_from_filled(rows: list[list[str]]) -> str:
    body_src = [_texts(row) for row in rows]
    width = max(len(row) for row in body_src)
    padded = [row + [""] * (width - len(row)) for row in body_src]
    return _df_to_markdown(padded[0], padded[1:])


def _outline_line(row: list[str]) -> str | None:
    filled = _filled(row)
    if not filled:
        return ""
    indent = min(filled[0][0], 6)
    pad = "  " * indent
    texts = [cell for _, cell in filled]
    if len(filled) == 1:
        if filled[0][0] == 0:
            return f"**{texts[0]}**"
        return pad + texts[0]
    if len(filled) >= 3:
        if filled[0][0] == 0 and not _is_numeric(texts[0]):
            return None
        if filled[0][0] > 0 and not _STEP.match(texts[0]):
            return f"{pad}{texts[0]}: {' '.join(texts[1:])}"
        return pad + " ".join(texts)
    if _STEP.match(texts[0]):
        return pad + " ".join(texts)
    return f"{pad}{texts[0]}: {texts[1]}"


def _outline_block(rows: list[list[str]]) -> str:
    chunks: list[str] = []
    index = 0
    while index < len(rows):
        run = _mini_run_len(rows, index)
        if run:
            chunks.append(_table_from_filled(rows[index : index + run]))
            index += run
            continue
        line = _outline_line(rows[index])
        if line is None:
            chunks.append(_table_from_filled(rows[index : index + 1]))
            index += 1
            continue
        if line:
            chunks.append(line)
        index += 1

    parts: list[str] = []
    for chunk in chunks:
        if not parts:
            parts.append(chunk)
            continue
        if chunk.startswith("**") or chunk.startswith("|") or parts[-1].startswith("|"):
            parts.append(f"\n{chunk}")
        else:
            parts.append(chunk)
    return "\n".join(parts).strip()


def _render_table(table: Table) -> list[str]:
    parts: list[str] = [f"## {table.name}"]
    if table.intro:
        parts.append(_render_intro(table.intro))
    if table.pairs:
        parts.append(_df_to_markdown(["key", "value"], [[k, v] for k, v in table.pairs]))
    if not table.rows:
        return parts
    if table.sparse:
        parts.append(_outline_block(table.rows))
        return parts
    for title, chunk in _split_banner_chunks(table.rows):
        if title:
            parts.append(f"**{title}**")
        if chunk:
            parts.append(_df_to_markdown(table.headers, chunk))
    return parts


def model_to_markdown(model: WorkbookModel, *, source_name: str) -> str:
    parts = [f"# {source_name}"]
    if model.meta:
        parts.append("## meta")
        parts.append(_df_to_markdown(["key", "value"], [[k, v] for k, v in model.meta.items()]))
    for table in model.tables:
        parts.extend(_render_table(table))
    return _MULTI_BLANK.sub("\n\n", "\n\n".join(part for part in parts if part).strip() + "\n")


def convert_path(
    path: str | Path,
    *,
    sheets: list[str] | None = None,
    markup: bool = True,
    raw: bool = False,
    keep_xlsx: str | Path | None = None,
    meta: str | Path | None = None,
    meta_info: dict | None = None,
) -> tuple[str, WorkbookModel | None]:
    path = Path(path)
    if raw:
        markdown = MarkItDown().convert(str(path)).markdown
        if meta:
            _write_meta(path, None, sheets=sheets, meta=meta, meta_info=meta_info)
        return markdown, None

    wb = load_workbook(path, data_only=True)
    try:
        model = process_workbook(wb, markup=markup, sheets=sheets)
        if meta:
            _write_meta(path, model, sheets=sheets, meta=meta, meta_info=meta_info, wb=wb)
    finally:
        wb.close()

    if keep_xlsx:
        write_tables_xlsx(model, str(keep_xlsx))
    return model_to_markdown(model, source_name=path.name), model


def _write_meta(
    path: Path,
    model: WorkbookModel | None,
    *,
    sheets: list[str] | None,
    meta: str | Path,
    meta_info: dict | None,
    wb=None,
) -> None:
    from .meta import build_meta, dump_meta

    owns_book = wb is None
    if owns_book:
        wb = load_workbook(path, data_only=True)
    try:
        document = build_meta(wb, model, source_name=path.name, sheets=sheets)
    finally:
        if owns_book:
            wb.close()
    dump_meta(document, meta)
    if meta_info is not None:
        meta_info["marks"] = len(document["marks"])
        meta_info["merges"] = len(document["merges"])


def format_stats(model: WorkbookModel | None, *, raw: bool) -> str:
    if raw or model is None:
        return "xlsx2md: raw markitdown (no preprocess)"
    lines = []
    for table in model.tables:
        shape = "outline" if table.sparse else f"{len(table.rows)}×{len(table.headers)}"
        lines.append(f"xlsx2md: {table.name}  {table.in_cells} cells → {shape}")
    lines.append(f"xlsx2md: {len(model.tables)} table(s)")
    return "\n".join(lines)
